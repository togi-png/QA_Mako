"""Synthetic proofs and sources for tests. No real campaign data."""
import io
from email.message import EmailMessage

import openpyxl

from qa import links as L

ACT, LOC = "A900001", "en_US"
CTA_URL = "https://www.example-adobe.test/products/demo.html?promoid=ABC&mv=email&trackingid=TAG00001"
CTA2_URL = "https://www.example-adobe.test/products/other.html?mv=email&trackingid=TAG00002"


def proof_html(body="Make something great today with the new demo app.", cta_text="Get started",
               extra="", optout=True, mirror=True, img_lang="en", merge_field=False, socials=None):
    socials = socials if socials is not None else ["Facebook", "Instagram"]
    social_html = "".join(
        f'<a href="https://t3.mail.adobe.com/r/?id=hSOC{n}" alias="{s}"><img alt="{s}" '
        f'src="https://landing.adobe.com/dam/global/images/social/{s.lower()}.png"></a>' for n, s in enumerate(socials))
    return f"""<html><head><style>.x{{color:red}}</style></head><body>
<div style="display:none; max-height:0px; overflow:hidden;">Preheader text here. &nbsp;&zwnj;&nbsp;&zwnj;</div>
<div id="cqa_technical" style="display:none;"><span id="cqa_activityid">{ACT}</span><span id="cqa_language">{LOC}</span></div>
<table><tr><td><h1>Big headline.</h1></td></tr>
<tr><td><p>{body}{' Hi <%= recipient.firstName %>' if merge_field else ''}</p></td></tr>
<tr><td><img alt="Hero" src="https://landing.adobe.com/dam/2026/images/1/demo.{img_lang}.1200x800.jpg"></td></tr>
<tr><td><a class="button" href="https://t3.mail.adobe.com/r/?id=hCTA1">{cta_text}</a></td></tr>
<tr><td><a href="https://t3.mail.adobe.com/r/?id=hCTA2">Learn more</a></td></tr>
{extra}
<tr><td>{social_html}</td></tr>
<tr><td><a href="https://t3.mail.adobe.com/r/?id=hFT1" alias="FOOTER - Terms">Terms of Use</a>
<a href="https://t3.mail.adobe.com/r/?id=hFT2" alias="FOOTER - Privacy Policy">Privacy Policy</a></td></tr>
{'<tr><td><a href="https://t3.mail.adobe.com/r/?id=hOPT" _type="optout" alias="Unsubscribe">unsubscribe</a></td></tr>' if optout else ''}
{'<tr><td><a href="https://t3.mail.adobe.com/r/?id=hMIR" _type="mirrorPage" alias="Read Online">Read online</a></td></tr>' if mirror else ''}
</table><img src="https://t3.mail.adobe.com/r/?id=hPIXEL,1" width="1" height="1"></body></html>"""


def proof_eml(subject="Make it yours", html=None, sender="Adobe Demo <mail@mail.adobe.com>", tag=True, **kw):
    m = EmailMessage()
    m["Subject"] = (f"[{ACT} {LOC} 123_20260901_US_EM1 Proof 1 multipart/alternative] " if tag else "") + subject
    m["From"] = sender
    m["To"] = "qa@example.test"
    m["List-Unsubscribe"] = "<https://unsubscribe.example.test/x>"
    m.set_content("plain text version")
    m.add_alternative(html or proof_html(**kw), subtype="html")
    return bytes(m)


def inline_forward(eml: bytes):
    import email
    from email import policy
    orig = email.message_from_bytes(eml, policy=policy.default)
    html = orig.get_body(("html",)).get_content()
    block = ('<div>Please QA these, thanks!</div><hr><div id="divRplyFwdMsg"><b>From:</b> Adobe Demo '
             '&lt;mail@mail.adobe.com&gt;<br><b>Sent:</b> Monday<br><b>To:</b> QA<br><b>Subject:</b> '
             + str(orig["Subject"]) + "</div>")
    html = html.replace("<body>", "<body>" + block, 1)
    m = EmailMessage()
    m["Subject"] = "FW: " + str(orig["Subject"])
    m["From"] = "Forwarder <someone@example.test>"
    m.set_content("see below")
    m.add_alternative(html, subtype="html")
    return bytes(m)


def container(*emls, attach_cgen: bytes = None):
    import email
    from email import policy
    m = EmailMessage()
    m["Subject"] = "FW: Campaign proofs"
    m["From"] = "Forwarder <someone@example.test>"
    m.set_content("Proofs attached")
    for e in emls:
        m.add_attachment(email.message_from_bytes(e, policy=policy.default))
    if attach_cgen:
        m.add_attachment(attach_cgen, maintype="application", subtype="vnd.ms-excel", filename="cgen.xlsx")
    return bytes(m)


def cgen_xlsx(headers=None, rows=None, banner=True):
    headers = headers or ["program id", "program name", "activity id", "activity name", "media type", "channel",
                          "tag id/code", "tag name", "tag purpose", "tag URL", "tag full url", "creative file name",
                          "subj line", "preheader", "deploy date", "From Name", "From Address"]
    rows = rows if rows is not None else [
        ["P1", "Demo program", ACT, "Email 1 Demo en_US", "Email", "Internal", "TAG00001", "Get started", "Join",
         CTA_URL.split("&trackingid")[0], CTA_URL, "123.en.demo", "Make it yours", "Preheader text here.",
         "9/1/2026", "Adobe Demo", "<mail@mail.adobe.com>"],
        ["P1", "Demo program", ACT, "Email 1 Demo en_US", "Email", "Internal", "TAG00002", "Learn more", "Join",
         CTA2_URL.split("&trackingid")[0], CTA2_URL, "123.en.demo", "Make it yours", "Preheader text here.",
         "9/1/2026", "Adobe Demo", "<mail@mail.adobe.com>"],
    ]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    if banner:
        ws.append(["Campaign Name: Demo"] + [None] * (len(headers) - 1))
    for r in rows:
        ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def reference_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Footer"
    ws.append(["locale", "link label", "accepted URL", "match type"])
    ws.append([LOC, "FOOTER - Terms", "https://www.adobe.com/legal/terms.html", "domain + path"])
    ws.append([LOC, "FOOTER - Privacy Policy", "https://www.adobe.com/privacy.html", "domain + path"])
    s = wb.create_sheet("Social")
    s.append(["program", "locale", "link label", "accepted URL", "match type"])
    for prog, handle in (("Creative Cloud", "adobecreativecloud"), ("MAX", "adobemax")):
        s.append([prog, LOC, "Facebook", f"https://www.facebook.com/{handle}", "domain + path"])
        s.append([prog, LOC, "Instagram", f"https://www.instagram.com/{handle}", "domain + path"])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


DEFAULT_MAP = {
    "hCTA1": CTA_URL, "hCTA2": CTA2_URL,
    "hSOC0": "https://www.facebook.com/adobecreativecloud", "hSOC1": "https://www.instagram.com/adobecreativecloud",
    "hFT1": "https://www.adobe.com/legal/terms.html", "hFT2": "https://www.adobe.com/privacy.html",
    "hMIR": "MIRROR",
}


class FakeResolver(L.Resolver):
    """Maps tracking links to landing URLs without any network access. Records every URL requested."""

    def __init__(self, mapping=None, mirror_html=None, down=False):
        super().__init__(enabled=True)
        self.mapping = dict(DEFAULT_MAP, **(mapping or {}))
        self.mirror_html = mirror_html
        self.down = down
        self.requested = []

    def resolve(self, url, want_body=False, _depth=0):
        self.requested.append(url)
        if self.down:
            return {"final_url": "", "status": None, "error": "Could not reach (ConnectionError). Check internet/VPN",
                    "chain": [], "body": ""}
        if "landing.adobe.com" in url:
            return {"final_url": url, "status": 200, "error": "", "chain": [], "body": ""}
        key = url.rsplit("=", 1)[-1].split(",")[0]
        dest = self.mapping.get(key)
        if dest == "MIRROR":
            return {"final_url": "https://t3.mail.adobe.com/m/?id=mirror", "status": 200, "error": "", "chain": [],
                    "body": self.mirror_html or ""}
        if dest is None:
            return {"final_url": "", "status": None, "error": "Could not reach (test)", "chain": [], "body": ""}
        return {"final_url": dest, "status": 200, "error": "", "chain": [url], "body": ""}
